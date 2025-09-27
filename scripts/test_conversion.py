#!/usr/bin/env python3
"""
Test script for Excel Windows conversion
Creates a sample Excel file and tests the conversion process
"""

import os
import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

def create_sample_excel():
    """Create a sample Excel file for testing"""
    print("🔧 Creating sample Excel file for testing...")

    # Create workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Test Sheet"

    # Add test data with formatting
    headers = ['ID', 'Name', 'Value', 'Date', 'Difference']
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')

    # Add sample data
    test_data = [
        [1, 'Item A', 100.50, '2024-01-01', -5.25],
        [2, 'Item B', 200.75, '2024-01-02', 3.75],
        [3, 'Item C', 150.00, '2024-01-03', 0.00],
        [4, 'Item D', 300.25, '2024-01-04', -2.50],
        [5, 'Item E', 250.00, '2024-01-05', 7.80]
    ]

    for row_idx, row_data in enumerate(test_data, 2):
        for col_idx, value in enumerate(row_data, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)

            # Apply borders
            thin_border = Border(
                left=Side(style='thin'), right=Side(style='thin'),
                top=Side(style='thin'), bottom=Side(style='thin')
            )
            cell.border = thin_border

            # Apply conditional formatting to difference column
            if col_idx == 5:  # Difference column
                if value < 0:
                    cell.font = Font(color='006100', bold=True)
                    cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                elif value > 0:
                    cell.font = Font(color='9C0006', bold=True)
                    cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

                cell.alignment = Alignment(horizontal='right')

    # Auto-fit columns
    for column in ws.columns:
        max_length = 0
        column_letter = column[0].column_letter
        for cell in column:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = (max_length + 2)
        ws.column_dimensions[column_letter].width = adjusted_width

    # Create file directory if it doesn't exist
    os.makedirs('file', exist_ok=True)

    # Save the file
    sample_file = 'file/test_sample.xlsx'
    wb.save(sample_file)
    print(f"✅ Sample file created: {sample_file}")

    return sample_file

def test_conversion():
    """Test the conversion process"""
    print("\n🧪 Testing conversion process...")

    # Import the conversion function
    sys.path.append('scripts')
    from convert_excel_windows import convert_excel_for_windows, create_csv_backup

    # Test file paths
    input_file = 'file/test_sample.xlsx'
    output_file = 'file_windows/test_sample_windows.xlsx'

    # Create output directory
    os.makedirs('file_windows', exist_ok=True)

    # Test conversion
    try:
        success = convert_excel_for_windows(input_file, output_file)
        if success:
            print("✅ Conversion test passed!")

            # Test CSV backup
            csv_success = create_csv_backup(input_file, 'file_windows')
            if csv_success:
                print("✅ CSV backup test passed!")

            # Verify files exist
            if os.path.exists(output_file):
                print(f"✅ Output file exists: {output_file}")

                # Check file size
                file_size = os.path.getsize(output_file)
                print(f"📊 Output file size: {file_size} bytes")

                return True
        else:
            print("❌ Conversion test failed!")
            return False

    except Exception as e:
        print(f"❌ Test failed with error: {e}")
        return False

def cleanup_test_files():
    """Clean up test files"""
    print("\n🧹 Cleaning up test files...")

    test_files = [
        'file/test_sample.xlsx',
        'file_windows/test_sample_windows.xlsx',
        'file_windows/test_sample_Test Sheet.csv'
    ]

    for file_path in test_files:
        if os.path.exists(file_path):
            os.remove(file_path)
            print(f"🗑️ Removed: {file_path}")

    # Remove directories if empty
    if os.path.exists('file') and not os.listdir('file'):
        os.rmdir('file')
        print("🗑️ Removed empty file/ directory")

    if os.path.exists('file_windows') and not os.listdir('file_windows'):
        os.rmdir('file_windows')
        print("🗑️ Removed empty file_windows/ directory")

def main():
    """Main test function"""
    print("🚀 Starting Excel conversion test...")

    try:
        # Create sample file
        sample_file = create_sample_excel()

        # Test conversion
        success = test_conversion()

        if success:
            print("\n🎉 All tests passed!")
            print("✅ Excel Windows conversion is working correctly")
        else:
            print("\n❌ Some tests failed!")
            print("🔧 Please check the conversion script")

        return 0 if success else 1

    except Exception as e:
        print(f"❌ Test suite failed: {e}")
        return 1

    finally:
        # Always cleanup
        cleanup_test_files()

if __name__ == "__main__":
    sys.exit(main())